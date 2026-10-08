"""Shared inputs of the SnapHiC / imputation benchmarks (``snaphic_*.py``, ``impute_*.py``).

* Bonev et al. 2017 mESC in situ Hi-C (4DN ``4DNFIC21MG3U.mcool``, registered as ``bonev2017_mesc_4dn``,
  12.6 GB): the 10 kb level of the autosomes is range-read over HTTP (fsspec + h5py; the ``indexes`` give the
  rows of each chromosome) and its band within ``max_dist`` is written to a local ``.cool`` with the file's
  ICE weights (``weight``, cooler balance, ``ignore_diags=2``) — the bulk map of the same cell type.
* SnapHiC's source data (Yu et al. 2021, Nat Methods 18:1056; Springer ESM, proposed registry entry
  ``yu2021_snaphic``): MOESM4 = SnapHiC loops from N = 10 ... 742 mES cells (six random subsets per N), MOESM6
  = the mESC reference lists (HiCCUPS on Bonev 2017, MAPS on H3K4me3 PLAC-seq, cohesin and H3K27ac HiChIP;
  ``*_filter``: 100 kb - 1 Mb, autosomes, mappable).
* Overlap rule of the paper: a loop (i, j) overlaps a loop (m, n) of another set when max(|i - m|, |j - n|)
  <= 20 kb (``uchrom.strc.loop.match_loops(tol=20_000)``).
"""

from __future__ import annotations

import json
import os
import resource
import sys
import time
from pathlib import Path
from typing import Dict, Optional, Sequence

import numpy as np
import pandas as pd

BONEV = ("https://4dn-open-data-public.s3.amazonaws.com/fourfront-webprod/wfoutput/"
         "2b448d99-d3fb-4bb9-a126-386ccb99a6fd/4DNFIC21MG3U.mcool")
MM10_AUTOSOMES = [f"chr{k}" for k in range(1, 20)]
TOL = 20_000


def data_dir() -> Path:
    from uchrom.datasets import data_dir as dd

    return dd()


def peak_rss_gb() -> float:
    r = resource.getrusage(resource.RUSAGE_SELF).ru_maxrss
    return r / 2**30 if sys.platform == "darwin" else r / 2**20


def load_avg() -> str:
    return " ".join(f"{x:.1f}" for x in os.getloadavg())


# ------------------------------------------------------------------------------------------------ Bonev 2017

def bonev_band(out: Path, *, chroms: Sequence[str] = MM10_AUTOSOMES, res: int = 10_000,
               max_dist: int = 2_600_000, url: str = BONEV, chunk: int = 20_000_000) -> Path:
    """The band (``j - i <= max_dist``) of the Bonev 2017 mESC map at ``res`` for ``chroms`` — range-read from
    the 4DN mcool over HTTP, or read from a local copy (``url`` a path) — into a ``.cool`` (built once)."""
    import cooler
    import fsspec
    import h5py

    out = Path(out)
    if out.exists():
        return out
    t0 = time.time()
    if "://" in str(url):
        fh = fsspec.open(url, block_size=2 ** 22, cache_type="blockcache").open()
        g = h5py.File(fh, "r")[f"resolutions/{res}"]
    else:                                                     # a local mcool (Sherlock: the whole file)
        g = h5py.File(str(url), "r")[f"resolutions/{res}"]
    names = [c.decode() if isinstance(c, bytes) else str(c) for c in g["chroms/name"][:]]
    lengths = g["chroms/length"][:]
    coff = g["indexes/chrom_offset"][:]
    b1off = g["indexes/bin1_offset"]
    weight_all = g["bins/weight"][:]
    wattrs = dict(g["bins/weight"].attrs)
    sizes = {c: int(lengths[names.index(c)]) for c in chroms}
    from uchrom.io.pairs_maps import genome_bins

    bins = genome_bins(sizes, res)
    w = []
    for c in chroms:
        k = names.index(c)
        w.append(weight_all[coff[k]:coff[k + 1]])
    bins["weight"] = np.concatenate(w)
    D = max_dist // res
    new_first = dict(zip(chroms, np.r_[0, np.cumsum([coff[names.index(c) + 1] - coff[names.index(c)]
                                                     for c in chroms])[:-1]].tolist()))

    def pixels():
        for c in chroms:
            k = names.index(c)
            lo, hi = int(coff[k]), int(coff[k + 1])
            rows = np.asarray(b1off[lo:hi + 1], np.int64)          # pixel offsets of the chromosome's rows
            kept, r0 = 0, 0
            while r0 < hi - lo:                                       # whole rows per read
                r1 = max(r0 + 1, int(np.searchsorted(rows, rows[r0] + chunk, side="right")) - 1)
                r1 = min(r1, hi - lo)
                a, b = int(rows[r0]), int(rows[r1])
                r0 = r1
                if b <= a:
                    continue
                b1 = g["pixels/bin1_id"][a:b]
                b2 = g["pixels/bin2_id"][a:b]
                cnt = g["pixels/count"][a:b]
                sel = (b2 < hi) & (b2 - b1 <= D)
                px = pd.DataFrame({"bin1_id": b1[sel] - lo + new_first[c], "bin2_id": b2[sel] - lo + new_first[c],
                                   "count": cnt[sel]})
                # repeated pixels (cooler refused the file's band as it is): summed
                px = px.groupby(["bin1_id", "bin2_id"], as_index=False, sort=True)["count"].sum()
                kept += len(px)
                yield px
            print(f"[bonev] {c}: {rows[-1] - rows[0]:,} pixels read, {kept:,} in the band ({time.time() - t0:.0f} s)",
                  flush=True)

    tmp = out.with_name(f".{out.name}.{os.getpid()}.part")
    out.parent.mkdir(parents=True, exist_ok=True)
    cooler.create_cooler(str(tmp), bins, pixels(), ordered=True, assembly="mm10",
                         metadata={"source": f"{url} resolutions/{res}, band <= {max_dist} bp, weights as stored"})
    with h5py.File(tmp, "r+") as h5:
        h5["bins/weight"].attrs.update({k: v for k, v in wattrs.items()})
    os.replace(tmp, out)
    return out


# ------------------------------------------------------------------------------------------- SnapHiC's data

def yu2021_dir() -> Path:
    return data_dir() / "yu2021_snaphic"


def published_sheets() -> list:
    """The sheet names of SnapHiC's source data (the TSV copies when there, else the workbooks)."""
    t = yu2021_dir() / "tsv"
    if t.exists():
        return sorted(p.stem for p in t.glob("*.tsv"))
    import openpyxl

    names = []
    for f in ("41592_2021_1231_MOESM4_ESM.xlsx", "41592_2021_1231_MOESM6_ESM.xlsx"):
        names += openpyxl.load_workbook(yu2021_dir() / f, read_only=True).sheetnames
    return names


def snaphic_published(sheet: str, *, path: Optional[Path] = None) -> pd.DataFrame:
    """A sheet of SnapHiC's source data as a loop table ``chrom1, start1, end1, chrom2, start2, end2``:
    ``"Permu0_742"`` etc. (MOESM4: SnapHiC's loops) or ``"Bulk_HiC_filter"`` etc. (MOESM6: references).  Read
    from ``yu2021_snaphic/tsv/<sheet>.tsv`` when present (the sheets written out once with openpyxl), else from
    the workbook."""
    tsv = yu2021_dir() / "tsv" / f"{sheet}.tsv"
    if path is None and tsv.exists():
        t = pd.read_csv(tsv, sep="\t")
    else:
        import openpyxl

        if path is None:
            path = yu2021_dir() / ("41592_2021_1231_MOESM4_ESM.xlsx" if sheet.startswith("Permu")
                                   else "41592_2021_1231_MOESM6_ESM.xlsx")
        wb = openpyxl.load_workbook(path, read_only=True)
        rows = list(wb[sheet].iter_rows(values_only=True))
        t = pd.DataFrame(rows[1:], columns=[str(c) for c in rows[0]])
    out = pd.DataFrame({"chrom1": "chr" + t["chr1"].astype(str), "start1": t["x1"].astype(np.int64),
                        "end1": t["x2"].astype(np.int64), "chrom2": "chr" + t["chr2"].astype(str),
                        "start2": t["y1"].astype(np.int64), "end2": t["y2"].astype(np.int64)})
    return out


def reference_union(*, filtered: bool = True) -> pd.DataFrame:
    """SnapHiC's combined mESC reference: HiCCUPS (bulk Hi-C) + MAPS (H3K4me3 PLAC-seq, cohesin and H3K27ac
    HiChIP), the ``*_filter`` sheets (100 kb - 1 Mb, autosomes, mappable)."""
    suf = "filter" if filtered else "all"
    parts = [snaphic_published(f"{s}_{suf}") for s in ("Bulk_HiC", "H3K4me3_PLACseq", "cohesin_HiChIP",
                                                      "H3K27ac_HiChIP")]
    return pd.concat(parts, ignore_index=True).drop_duplicates(["chrom1", "start1", "start2"]).reset_index(drop=True)


# ------------------------------------------------------------------------------------------------ overlaps

def _index(table: pd.DataFrame) -> Dict[str, np.ndarray]:
    out = {}
    for c, g in table.groupby(table["chrom1"].astype(str)):
        out[c] = g[["start1", "start2"]].to_numpy(np.int64)
    return out


def overlap(a: pd.DataFrame, b: pd.DataFrame, tol: int = TOL) -> np.ndarray:
    """For every loop of ``a``: is a loop of ``b`` on its chromosome with both anchors within ``tol``?"""
    from scipy.spatial import cKDTree

    hit = np.zeros(len(a), bool)
    bi = _index(b)
    a = a.reset_index(drop=True)
    for c, g in a.groupby(a["chrom1"].astype(str)):
        if c not in bi or not len(bi[c]):
            continue
        tree = cKDTree(bi[c].astype(np.float64))
        d, _ = tree.query(g[["start1", "start2"]].to_numpy(np.float64), p=np.inf)
        hit[g.index.to_numpy()] = d <= tol
    return hit


def precision_recall(calls: pd.DataFrame, reference: pd.DataFrame, tol: int = TOL, *,
                     precision_reference: Optional[pd.DataFrame] = None) -> Dict[str, float]:
    """SnapHiC's measures: precision = calls overlapping the (precision) reference / calls; recall =
    reference loops overlapped by a call / reference; F1."""
    pref = reference if precision_reference is None else precision_reference
    n_calls, n_ref = len(calls), len(reference)
    p = float(overlap(calls, pref, tol).mean()) if n_calls else float("nan")
    r = float(overlap(reference, calls, tol).mean()) if n_ref else float("nan")
    f1 = 2 * p * r / (p + r) if n_calls and n_ref and (p + r) > 0 else float("nan")
    return {"n_calls": int(n_calls), "n_reference": int(n_ref), "precision": p, "recall": r, "f1": f1}


def in_range(t: pd.DataFrame, lo: int = 100_000, hi: int = 1_000_000, chroms: Optional[Sequence[str]] = None) -> pd.DataFrame:
    sep = t["start2"] - t["start1"]
    keep = (sep >= lo) & (sep <= hi)
    if chroms is not None:
        keep &= t["chrom1"].astype(str).isin(list(chroms))
    return t[keep].reset_index(drop=True)


def write_json(obj, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(obj, indent=1, default=lambda x: x.item() if hasattr(x, "item") else str(x)))
    print(f"-> {path}")
